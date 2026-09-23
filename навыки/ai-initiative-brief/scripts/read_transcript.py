#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
read_transcript.py — извлекает простой текст из файла транскрипта интервью
и печатает его в stdout, чтобы навык мог прочитать содержимое.

Поддержка: .txt, .md, .docx (python-docx), .pdf (pdfplumber или pypdf, если есть).

Использование:
  python3 read_transcript.py "<путь к файлу>"
"""

import sys
from pathlib import Path


def read_txt(path):
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def read_docx(path):
    try:
        from docx import Document
    except ImportError:
        sys.exit(
            "Для .docx нужен python-docx: "
            "<python-из-venv> -m pip install python-docx (см. README.md)"
        )
    doc = Document(str(path))
    out = []
    for p in doc.paragraphs:
        if p.text.strip():
            out.append(p.text)
    # Текст из таблиц (иногда транскрипты — двухколоночные: спикер | реплика)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append("\t".join(cells))
    return "\n".join(out)


def read_pdf(path):
    # Сначала pdfplumber (точнее), затем pypdf
    try:
        import pdfplumber

        chunks = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                t = page.extract_text() or ""
                if t.strip():
                    chunks.append(t)
        return "\n".join(chunks)
    except ImportError:
        pass
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except ImportError:
        sys.exit(
            "Для .pdf нужен pdfplumber или pypdf "
            "(<python-из-venv> -m pip install pypdf; см. README.md). "
            "Также можно продолжить интервью в чате."
        )


def main():
    if len(sys.argv) < 2:
        sys.exit("Использование: python3 read_transcript.py <путь к файлу>")
    path = Path(sys.argv[1])
    if not path.exists():
        sys.exit(f"Файл не найден: {path}")

    ext = path.suffix.lower()
    if ext in (".txt", ".md", ".markdown"):
        text = read_txt(path)
    elif ext == ".docx":
        text = read_docx(path)
    elif ext == ".pdf":
        text = read_pdf(path)
    else:
        sys.exit(f"Неподдерживаемый формат: {ext}. Используйте .txt/.md/.docx/.pdf")

    if not text.strip():
        sys.exit("Текст не найден. Для скана нужно OCR; можно предоставить текст или пройти интервью в чате.")
    sys.stdout.write(text)


if __name__ == "__main__":
    main()
