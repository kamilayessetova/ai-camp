#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_brief.py — собирает Word-документ (.docx) «Бриф на ИИ-инициативу»
из заполненного markdown-файла, написанного по структуре assets/template.md.

Сохраняет курсивные подсказки-инструкции шаблона (строки вида *...*) как
серые курсивные памятки под заголовками — это требование владельца шаблона.

Поддерживаемое подмножество markdown (ровно то, что есть в шаблоне):
  # Заголовок            -> Title (главный заголовок брифа)
  ## Раздел*             -> Heading 1 (звёздочка обязательности сохраняется)
  ### Подраздел          -> Heading 2
  *курсивная подсказка*  -> серый курсив (памятка)
  **Метка:** текст       -> абзац с жирным фрагментом (инлайн **...**)
  1. пункт               -> нумерованный список (номер берётся из исходника)
  - пункт                -> маркированный список
  > ❓ текст             -> видимый плейсхолдер «требует уточнения» (красный)
  пустая строка          -> пропускается
  прочее                 -> обычный абзац (с поддержкой инлайн **жирного**)

Использование:
  python3 build_brief.py "brief_filled.md" "Бриф на ИИ-инициативу — Название.docx"
Если второй аргумент не задан, имя берётся из заголовка # внутри файла.
"""

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    sys.exit(
        "Не найден python-docx. Установите: "
        "<python-из-venv> -m pip install python-docx (см. README.md)"
    )

HINT_GRAY = RGBColor(0x70, 0x70, 0x70)
PLACEHOLDER_RED = RGBColor(0xC0, 0x00, 0x00)

NUM_RE = re.compile(r"^(\d+)\.\s+(.*)$")
BULLET_RE = re.compile(r"^[-•]\s+(.*)$")


def add_inline(paragraph, text, *, italic=False, color=None, bold_all=False):
    """Добавляет в абзац runs, разбирая инлайн **жирный**."""
    parts = text.split("**")
    for i, part in enumerate(parts):
        if part == "":
            continue
        run = paragraph.add_run(part)
        run.bold = bold_all or (i % 2 == 1)  # нечётные сегменты — внутри **...**
        run.italic = italic
        if color is not None:
            run.font.color.rgb = color


def is_hint(line):
    """Строка-подсказка целиком в одиночных звёздочках: *...* (не **...**)."""
    s = line.strip()
    return (
        len(s) >= 2
        and s.startswith("*")
        and s.endswith("*")
        and not s.startswith("**")
        and not s.endswith("**")
    )


def build(md_path, out_path=None):
    md_path = Path(md_path)
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    # Базовый шрифт
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title_text = None

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()

        if stripped == "":
            continue

        # Заголовки
        if stripped.startswith("# "):
            title_text = stripped[2:].strip()
            doc.add_heading(title_text, level=0)
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
            continue

        # Плейсхолдер «требует уточнения»
        if stripped.startswith(">"):
            content = stripped.lstrip(">").strip()
            p = doc.add_paragraph()
            add_inline(p, content, italic=True, color=PLACEHOLDER_RED, bold_all=True)
            continue

        # Курсивная подсказка-памятка
        if is_hint(stripped):
            p = doc.add_paragraph()
            add_inline(p, stripped[1:-1], italic=True, color=HINT_GRAY)
            continue

        # Нумерованный список — номер сохраняем из исходника
        m = NUM_RE.match(stripped)
        if m:
            number, content = m.group(1), m.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(f"{number}. ")
            run.bold = False
            add_inline(p, content)
            continue

        # Маркированный список
        m = BULLET_RE.match(stripped)
        if m:
            content = m.group(1)
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, content)
            continue

        # Обычный абзац
        p = doc.add_paragraph()
        add_inline(p, stripped)

    # Имя файла
    if out_path is None:
        base = title_text or md_path.stem
        out_path = md_path.with_name(f"{base}.docx")
    out_path = Path(out_path)
    doc.save(str(out_path))
    return out_path


def main():
    if len(sys.argv) < 2:
        sys.exit("Использование: python3 build_brief.py <filled.md> [output.docx]")
    md_path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else None
    saved = build(md_path, out_path)
    print(f"Готово: {saved}")


if __name__ == "__main__":
    main()
